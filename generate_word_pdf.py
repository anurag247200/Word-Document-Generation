import docx
from docx2pdf import convert

# Step 1: Load the Word template
doc = docx.Document('template.docx')  # Ensure this file exists in the same directory as your script

# Step 2: Dictionary with placeholders and their values
replacements = {
    '{{name}}': 'Anurag Singh',
    '{{address}}': '123 Main Street, Kolkata',
    '{{date}}': 'September 29, 2024'
}

# Step 3: Function to replace placeholders in paragraphs
def replace_placeholders(paragraphs, replacements):
    for paragraph in paragraphs:
        for key, value in replacements.items():
            if key in paragraph.text:
                paragraph.text = paragraph.text.replace(key, value)

# Step 4: Replace placeholders in the document's paragraphs
replace_placeholders(doc.paragraphs, replacements)

# Step 5: Save the modified Word document
output_docx = 'output.docx'
doc.save(output_docx)

print(f"Word document saved as {output_docx}")

# Step 6: Convert the Word document to PDF
output_pdf = 'output.pdf'
convert(output_docx, output_pdf)

print(f"PDF file saved as {output_pdf}")

